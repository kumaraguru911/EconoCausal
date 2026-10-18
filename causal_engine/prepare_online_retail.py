"""Stream and audit UCI Online Retail II transaction data."""
import argparse
import csv
import json
from collections import Counter
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

DEFAULT_ARCHIVE = Path("data/raw/online_retail_ii/online_retail_ii.zip")
DEFAULT_OUTPUT = Path("data/processed/online_retail_ii_clean.csv")
DEFAULT_AUDIT = Path("reports/online_retail_ii_audit.json")

OUTPUT_COLUMNS = [
    "Invoice", "StockCode", "Description", "Quantity",
    "InvoiceDate", "Price", "Customer ID", "Country",
    "LineRevenue", "SourceSheet",
]


def serialise(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return value


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_ARCHIVE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--audit", type=Path, default=DEFAULT_AUDIT)
    args = parser.parse_args()

    if not args.input.exists():
        raise FileNotFoundError(f"Dataset not found: {args.input}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.audit.parent.mkdir(parents=True, exist_ok=True)

    audit = {
        "input": str(args.input),
        "output": str(args.output),
        "sheets": {},
        "total_rows": 0,
        "written_rows": 0,
        "excluded_cancelled_invoice": 0,
        "excluded_nonpositive_quantity": 0,
        "excluded_nonpositive_price": 0,
        "excluded_missing_required_fields": 0,
        "missing_values": {},
        "notes": [
            "This is observational transaction data, not a randomized price experiment.",
            "The clean file excludes cancellation invoices and non-positive quantities/prices.",
            "Returns and cancellations are excluded from the primary sales file; audit counts are retained.",
            "LineRevenue is Quantity multiplied by recorded unit Price.",
        ],
    }
    missing_counts = Counter()

    with ZipFile(args.input) as archive:
        with archive.open("online_retail_II.xlsx") as workbook_bytes:
            workbook_data = BytesIO(workbook_bytes.read())

    workbook = load_workbook(workbook_data, read_only=True, data_only=True)

    try:
        with args.output.open("w", newline="", encoding="utf-8") as output_file:
            writer = csv.DictWriter(output_file, fieldnames=OUTPUT_COLUMNS)
            writer.writeheader()

            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    continue

                positions = {str(name).strip(): i for i, name in enumerate(headers)}
                required = ["Invoice", "StockCode", "Quantity", "InvoiceDate", "Price"]
                missing_headers = [name for name in required if name not in positions]
                if missing_headers:
                    raise ValueError(
                        f"Sheet {sheet.title!r} is missing columns: {missing_headers}"
                    )

                stats = Counter()
                for row in rows:
                    stats["rows"] += 1
                    audit["total_rows"] += 1
                    record = {
                        name: row[index] if index < len(row) else None
                        for name, index in positions.items()
                    }

                    for name, value in record.items():
                        if value is None or (isinstance(value, str) and not value.strip()):
                            missing_counts[name] += 1

                    required_values = [
                        record.get("Invoice"),
                        record.get("StockCode"),
                        record.get("Quantity"),
                        record.get("InvoiceDate"),
                        record.get("Price"),
                    ]
                    if any(value is None for value in required_values):
                        stats["missing_required"] += 1
                        audit["excluded_missing_required_fields"] += 1
                        continue

                    invoice = str(record["Invoice"]).strip()
                    if invoice.upper().startswith("C"):
                        stats["cancelled_invoice"] += 1
                        audit["excluded_cancelled_invoice"] += 1
                        continue

                    try:
                        quantity = float(record["Quantity"])
                        price = float(record["Price"])
                    except (TypeError, ValueError):
                        stats["invalid_numeric"] += 1
                        audit["excluded_missing_required_fields"] += 1
                        continue

                    if quantity <= 0:
                        stats["nonpositive_quantity"] += 1
                        audit["excluded_nonpositive_quantity"] += 1
                        continue

                    if price <= 0:
                        stats["nonpositive_price"] += 1
                        audit["excluded_nonpositive_price"] += 1
                        continue

                    output_record = {
                        "Invoice": invoice,
                        "StockCode": record.get("StockCode"),
                        "Description": record.get("Description"),
                        "Quantity": quantity,
                        "InvoiceDate": serialise(record["InvoiceDate"]),
                        "Price": price,
                        "Customer ID": record.get("Customer ID"),
                        "Country": record.get("Country"),
                        "LineRevenue": quantity * price,
                        "SourceSheet": sheet.title,
                    }
                    writer.writerow({
                        key: "" if value is None else value
                        for key, value in output_record.items()
                    })
                    stats["written"] += 1
                    audit["written_rows"] += 1

                audit["sheets"][sheet.title] = dict(stats)
                print(
                    f"{sheet.title}: rows={stats['rows']:,}, "
                    f"written={stats['written']:,}, "
                    f"cancelled={stats['cancelled_invoice']:,}, "
                    f"nonpositive_quantity={stats['nonpositive_quantity']:,}, "
                    f"nonpositive_price={stats['nonpositive_price']:,}"
                )
    finally:
        workbook.close()

    audit["missing_values"] = dict(missing_counts)
    with args.audit.open("w", encoding="utf-8") as audit_file:
        json.dump(audit, audit_file, indent=2, default=str)

    print(f"\nTotal input rows: {audit['total_rows']:,}")
    print(f"Clean rows written: {audit['written_rows']:,}")
    print(f"Clean CSV: {args.output}")
    print(f"Audit report: {args.audit}")


if __name__ == "__main__":
    main()
